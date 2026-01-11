import io
import json
import csv
import base64
from typing import Dict, List, Any, Optional
from datetime import datetime
from pathlib import Path
import structlog

from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt, Mm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from app.core.config import settings
from app.schemas.therapy import TherapyRecommendation

logger = structlog.get_logger(__name__)

# Logo path
LOGO_PATH = Path(__file__).parent.parent / "static" / "logo.png"
REPORT_DESCRIPTION = (
    "Theramuse Rx is research-driven music therapy assistant designed for personalised "
    "interventions for dementia, Down Syndrome, and ADHD populations. The system ingests "
    "rich intake assessments, maps them to generational nostalgia cues, and curates long-form "
    "content from YouTube or other music streaming platforms with rigorous filtering."
)


class ExportService:
    """
    Service for exporting patient data and recommendations in various formats.
    """

    @staticmethod
    def _safe_text(value: Any) -> str:
        """
        FPDF core fonts are latin-1 only. Replace unsupported characters so
        the export does not crash on emoji or extended glyphs.
        Uses improved character mapping for better Unicode handling.
        """
        if value is None:
            return ""

        text = str(value)
        # Common Unicode character replacements for better readability
        replacements = {
            '"': '"',
            '"': '"',
            ''': "'",
            ''': "'",
            '–': '-',
            '—': '--',
            '…': '...',
            '©': '(C)',
            '®': '(R)',
            '™': '(TM)',
            '°': ' deg ',
            '±': '+/-',
            '≤': '<=',
            '≥': '>=',
            '≈': '~',
            '≠': '!=',
            '→': '->',
            '←': '<-',
            '↑': '^',
            '↓': 'v',
        }

        for unicode_char, replacement in replacements.items():
            text = text.replace(unicode_char, replacement)

        # For any remaining unsupported characters, use latin-1 encoding with replacement
        return text.encode("latin-1", errors="replace").decode("latin-1")

    def __init__(self):
        self.export_dir = Path(settings.EXPORT_DIR)
        self.export_dir.mkdir(exist_ok=True)

    def generate_pdf_report(self, patient_info: Dict[str, Any],
                          recommendations: List[TherapyRecommendation],
                          big_five: Optional[Dict[str, float]] = None,
                          patient_summary: Optional[Dict[str, Any]] = None) -> bytes:
        """
        Generate PDF report for patient data and recommendations.
        """
        try:
            pdf = FPDF()
            pdf.add_page()

            # Add logo at the top if available
            if LOGO_PATH.exists():
                try:
                    # Calculate width to be approximately 50mm (about 2 inches)
                    # Height will be calculated automatically to maintain aspect ratio
                    logo_width = 50
                    # Center the logo horizontally
                    page_width = pdf.w
                    x_position = (page_width - logo_width) / 2
                    pdf.image(str(LOGO_PATH), x=x_position, y=10, w=logo_width)
                    pdf.ln(35)  # Move down after logo
                except Exception as e:
                    logger.warning(f"Failed to add logo to PDF: {e}")
                    pdf.ln(5)
            else:
                pdf.ln(5)

            pdf.set_font("Arial", "", 9)
            pdf.multi_cell(0, 5, self._safe_text(REPORT_DESCRIPTION))
            pdf.ln(6)

            pdf.set_font("Arial", "B", 16)

            # Title
            pdf.cell(0, 10, "TheraMuse Music Therapy Report", ln=True, align="C")
            pdf.ln(10)

            # Patient Information
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, "Client Information", ln=True)
            pdf.set_font("Arial", "", 10)

            for key, value in patient_info.items():
                if value is None:
                    continue
                normalized_key = key.replace(" ", "_").lower()
                if normalized_key in {"session_id", "sessionid"}:
                    continue
                label = key.replace('_', ' ').title().replace("Patient", "Client")
                pdf.cell(0, 6, self._safe_text(f"{label}: {value}"), ln=True)

            pdf.ln(10)

            # Big Five Scores (if available)
            if big_five:
                # Convert BigFiveScores object to dictionary if needed
                if hasattr(big_five, 'model_dump'):
                    big_five_dict = big_five.model_dump()
                elif hasattr(big_five, 'dict'):
                    big_five_dict = big_five.dict()
                else:
                    big_five_dict = big_five

                pdf.set_font("Arial", "B", 12)
                pdf.cell(0, 10, "Big Five Personality Scores", ln=True)
                pdf.set_font("Arial", "", 10)

                for trait, score in big_five_dict.items():
                    pdf.cell(0, 6, self._safe_text(f"{trait.title()}: {score:.2f}"), ln=True)
                pdf.ln(10)

            # Patient Summary (if available)
            if patient_summary:
                pdf.set_font("Arial", "B", 12)
                pdf.cell(0, 10, "Session Summary", ln=True)
                pdf.set_font("Arial", "", 10)

                for key, value in patient_summary.items():
                    if key in {"session_id", "sessionId"}:
                        continue
                    if isinstance(value, (str, int, float)):
                        pdf.cell(0, 6, self._safe_text(f"{key.replace('_', ' ').title()}: {value}"), ln=True)
                pdf.ln(10)

            # Recommendations
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, "Prescribed Songs", ln=True)
            pdf.set_font("Arial", "", 10)

            for i, rec in enumerate(recommendations, 1):
                pdf.cell(0, 8, self._safe_text(f"{i}. {rec.song_title}"), ln=True)
                if rec.artist:
                    pdf.cell(0, 6, self._safe_text(f"   Artist: {rec.artist}"), ln=True)
                if rec.genre:
                    pdf.cell(0, 6, self._safe_text(f"   Genre: {rec.genre}"), ln=True)
                if rec.year:
                    pdf.cell(0, 6, self._safe_text(f"   Year: {rec.year}"), ln=True)
                if rec.youtube_url:
                    pdf.cell(0, 6, self._safe_text(f"   YouTube: {rec.youtube_url}"), ln=True)
                if rec.spotify_url:
                    pdf.cell(0, 6, self._safe_text(f"   Spotify: {rec.spotify_url}"), ln=True)
                if rec.recommendation_score is not None:
                    pdf.cell(0, 6, self._safe_text(f"   Score: {rec.recommendation_score:.2f}"), ln=True)
                pdf.cell(0, 6, self._safe_text(f"   Recommended: {rec.created_at.strftime('%Y-%m-%d %H:%M')}"), ln=True)
                pdf.ln(3)

            # Footer
            pdf.set_font("Arial", "I", 8)
            pdf.cell(0, 10, f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True, align="C")
            pdf.cell(0, 5, "TheraMuse - Personalized Music Therapy System", ln=True, align="C")

            pdf_buffer = pdf.output(dest="S")
            if isinstance(pdf_buffer, bytes):
                return pdf_buffer
            elif isinstance(pdf_buffer, bytearray):
                return bytes(pdf_buffer)
            else:
                return str(pdf_buffer).encode("latin-1", errors="replace")

        except Exception as e:
            logger.error(f"Failed to generate PDF report: {e}")
            raise

    def generate_docx_report(self, patient_info: Dict[str, Any],
                           recommendations: List[TherapyRecommendation],
                           big_five: Optional[Dict[str, float]] = None,
                           patient_summary: Optional[Dict[str, Any]] = None) -> bytes:
        """
        Generate DOCX report for patient data and recommendations.
        """
        try:
            doc = Document()

            # Add logo at the top if available
            if LOGO_PATH.exists():
                try:
                    # Add logo centered at the top (approximately 2 inches wide)
                    header_para = doc.add_paragraph()
                    header_para.alignment = 1  # Center alignment
                    run = header_para.add_run()
                    run.add_picture(str(LOGO_PATH), width=Inches(2.0))
                except Exception as e:
                    logger.warning(f"Failed to add logo to DOCX: {e}")

            description_para = doc.add_paragraph()
            description_run = description_para.add_run(REPORT_DESCRIPTION)
            description_run.font.size = Pt(9)
            doc.add_paragraph()

            # Title
            title = doc.add_heading('TheraMuse Music Therapy Report', 0)
            title.alignment = 1  # Center alignment

            # Patient Information
            doc.add_heading('Client Information', level=1)
            client_info_items = []
            for key, value in patient_info.items():
                if value is None:
                    continue
                normalized_key = key.replace(" ", "_").lower()
                if normalized_key in {"session_id", "sessionid"}:
                    continue
                label = key.replace('_', ' ').title().replace("Patient", "Client")
                client_info_items.append((label, value))

            if not client_info_items:
                client_info_items.append(("Client Info", ""))

            patient_table = doc.add_table(rows=len(client_info_items), cols=2)
            patient_table.style = 'Table Grid'

            for i, (label, value) in enumerate(client_info_items):
                patient_table.cell(i, 0).text = str(label)
                patient_table.cell(i, 1).text = str(value)

            doc.add_paragraph()  # Add space

            # Big Five Scores
            if big_five:
                # Convert BigFiveScores object to dictionary if needed
                if hasattr(big_five, 'model_dump'):
                    big_five_dict = big_five.model_dump()
                elif hasattr(big_five, 'dict'):
                    big_five_dict = big_five.dict()
                else:
                    big_five_dict = big_five

                doc.add_heading('Big Five Personality Scores', level=1)
                big5_table = doc.add_table(rows=len(big_five_dict) + 1, cols=2)
                big5_table.style = 'Table Grid'

                big5_table.cell(0, 0).text = 'Trait'
                big5_table.cell(0, 1).text = 'Score'

                for i, (trait, score) in enumerate(big_five_dict.items(), 1):
                    big5_table.cell(i, 0).text = trait.title()
                    big5_table.cell(i, 1).text = f"{score:.3f}"

                doc.add_paragraph()  # Add space

            # Patient Summary
            if patient_summary:
                doc.add_heading('Session Summary', level=1)
                for key, value in patient_summary.items():
                    if key in {"session_id", "sessionId"}:
                        continue
                    if isinstance(value, (str, int, float)):
                        doc.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")

                doc.add_paragraph()  # Add space

            # Recommendations
            doc.add_heading('Prescribed Songs', level=1)

            for i, rec in enumerate(recommendations, 1):
                rec_heading = doc.add_heading(f'{i}. {rec.song_title}', level=2)
                rec_para = doc.add_paragraph()

                if rec.artist:
                    rec_para.add_run(f'Artist: {rec.artist}\n')
                if rec.genre:
                    rec_para.add_run(f'Genre: {rec.genre}\n')
                if rec.year:
                    rec_para.add_run(f'Year: {rec.year}\n')
                if rec.youtube_url:
                    rec_para.add_run(f'YouTube: {rec.youtube_url}\n')
                if rec.spotify_url:
                    rec_para.add_run(f'Spotify: {rec.spotify_url}\n')
                if rec.recommendation_score is not None:
                    rec_para.add_run(f'Recommendation Score: {rec.recommendation_score:.2f}\n')
                rec_para.add_run(f'Recommended: {rec.created_at.strftime("%Y-%m-%d %H:%M")}\n')

            # Footer
            doc.add_paragraph()
            footer_para = doc.add_paragraph()
            footer_para.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            footer_para.add_run("\nTheraMuse - Personalized Music Therapy System")
            footer_para.alignment = 1  # Center alignment

            # Save to bytes
            doc_bytes = io.BytesIO()
            doc.save(doc_bytes)
            return doc_bytes.getvalue()

        except Exception as e:
            logger.error(f"Failed to generate DOCX report: {e}")
            raise

    def generate_csv_export(self, patient_info: Dict[str, Any],
                          recommendations: List[TherapyRecommendation],
                          big_five: Optional[Dict[str, float]] = None,
                          patient_summary: Optional[Dict[str, Any]] = None) -> bytes:
        """
        Generate CSV export of patient data and recommendations.
        """
        try:
            output = io.StringIO()
            writer = csv.writer(output)

            # Header
            writer.writerow(['TheraMuse Export'])
            writer.writerow(['Generated:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
            writer.writerow([])

            # Patient Information
            writer.writerow(['Patient Information'])
            writer.writerow(['Field', 'Value'])
            for key, value in patient_info.items():
                if value is not None:
                    writer.writerow([key.replace('_', ' ').title(), str(value)])

            writer.writerow([])

            # Big Five Scores
            if big_five:
                # Convert BigFiveScores object to dictionary if needed
                if hasattr(big_five, 'model_dump'):
                    big_five_dict = big_five.model_dump()
                elif hasattr(big_five, 'dict'):
                    big_five_dict = big_five.dict()
                else:
                    big_five_dict = big_five

                writer.writerow(['Big Five Personality Scores'])
                writer.writerow(['Trait', 'Score'])
                for trait, score in big_five_dict.items():
                    writer.writerow([trait.title(), f"{score:.3f}"])
                writer.writerow([])

            # Recommendations
            writer.writerow(['Music Recommendations'])
            writer.writerow(['Rank', 'Title', 'Artist', 'Genre', 'Year', 'YouTube URL', 'Spotify URL', 'Score', 'Recommended Date'])

            for rec in recommendations:
                writer.writerow([
                    rec.rank or '',
                    rec.song_title,
                    rec.artist or '',
                    rec.genre or '',
                    rec.year or '',
                    rec.youtube_url or '',
                    rec.spotify_url or '',
                    f"{rec.recommendation_score:.3f}" if rec.recommendation_score is not None else '',
                    rec.created_at.strftime('%Y-%m-%d %H:%M')
                ])

            return output.getvalue().encode('utf-8')

        except Exception as e:
            logger.error(f"Failed to generate CSV export: {e}")
            raise

    def generate_json_export(self, patient_info: Dict[str, Any],
                           recommendations: List[TherapyRecommendation],
                           big_five: Optional[Dict[str, float]] = None,
                           patient_summary: Optional[Dict[str, Any]] = None) -> bytes:
        """
        Generate JSON export of patient data and recommendations.
        """
        try:
            export_data = {
                'export_info': {
                    'generated_at': datetime.now().isoformat(),
                    'system': 'TheraMuse',
                    'version': '1.0.0'
                },
                'patient_info': patient_info,
                'big_five_scores': big_five,
                'patient_summary': patient_summary,
                'recommendations': []
            }

            for rec in recommendations:
                rec_data = {
                    'rank': rec.rank,
                    'song_title': rec.song_title,
                    'artist': rec.artist,
                    'genre': rec.genre,
                    'year': rec.year,
                    'tempo': rec.tempo,
                    'valence': rec.valence,
                    'arousal': rec.arousal,
                    'youtube_url': rec.youtube_url,
                    'spotify_url': rec.spotify_url,
                    'recommendation_score': rec.recommendation_score,
                    'algorithm_used': rec.algorithm_used,
                    'created_at': rec.created_at.isoformat()
                }
                export_data['recommendations'].append(rec_data)

            return json.dumps(export_data, indent=2, ensure_ascii=False).encode('utf-8')

        except Exception as e:
            logger.error(f"Failed to generate JSON export: {e}")
            raise

    def export_data(self, export_format: str, patient_info: Dict[str, Any],
                   recommendations: List[TherapyRecommendation],
                   big_five: Optional[Dict[str, float]] = None,
                   patient_summary: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """
        Export data in the specified format.
        """
        try:
            export_format = export_format.lower()
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

            # Generate content based on format
            if export_format == 'pdf':
                content = self.generate_pdf_report(
                    patient_info, recommendations, big_five, patient_summary
                )
                filename = f"theramuse_report_{timestamp}.pdf"
                content_type = "application/pdf"

            elif export_format == 'docx':
                content = self.generate_docx_report(
                    patient_info, recommendations, big_five, patient_summary
                )
                filename = f"theramuse_report_{timestamp}.docx"
                content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

            elif export_format == 'csv':
                content = self.generate_csv_export(
                    patient_info, recommendations, big_five, patient_summary
                )
                filename = f"theramuse_export_{timestamp}.csv"
                content_type = "text/csv"

            elif export_format == 'json':
                content = self.generate_json_export(
                    patient_info, recommendations, big_five, patient_summary
                )
                filename = f"theramuse_export_{timestamp}.json"
                content_type = "application/json"

            else:
                raise ValueError(f"Unsupported export format: {export_format}")

            # Create response data
            result = {
                'filename': filename,
                'content_type': content_type,
                'file_size': len(content),
                'base64_data': base64.b64encode(content).decode('utf-8')
            }

            logger.info(
                f"Successfully exported data in {export_format} format",
                filename=filename,
                file_size=len(content)
            )

            return result

        except Exception as e:
            logger.error(f"Failed to export data: {e}")
            raise
