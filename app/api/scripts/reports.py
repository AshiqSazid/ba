from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any, Dict, Iterable, List, Tuple


def generate_docx_report(
    patient_info: Dict[str, Any],
    recommendations: Dict[str, Any],
    big5_scores: Dict[str, Any],
) -> bytes:
    from docx import Document  # type: ignore

    document = Document()
    document.add_heading('TheraMuse Recommendation Report', level=0)
    document.add_paragraph(f'Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}')

    document.add_heading('Patient Overview', level=1)
    patient_table = document.add_table(rows=1, cols=2)
    for idx, (label, value) in enumerate(_patient_rows(patient_info)):
        cells = patient_table.rows[0].cells if idx == 0 else patient_table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading('Big Five Scores', level=1)
    big5_table = document.add_table(rows=1, cols=2)
    header_cells = big5_table.rows[0].cells
    header_cells[0].text = 'Trait'
    header_cells[1].text = 'Score'
    for trait, score in _big5_rows(big5_scores):
        row_cells = big5_table.add_row().cells
        row_cells[0].text = trait
        row_cells[1].text = score

    document.add_heading('Recommendations', level=1)
    rec_rows = _recommendation_rows(recommendations)
    if rec_rows:
        rec_table = document.add_table(rows=1, cols=5)
        headers = rec_table.rows[0].cells
        headers[0].text = 'Category'
        headers[1].text = '#'
        headers[2].text = 'Title'
        headers[3].text = 'Channel'
        headers[4].text = 'Notes'
        for row in rec_rows:
            rec_cells = rec_table.add_row().cells
            rec_cells[0].text = row['category']
            rec_cells[1].text = str(row['rank'])
            rec_cells[2].text = row['title']
            rec_cells[3].text = row['channel']
            rec_cells[4].text = row['description']
    else:
        document.add_paragraph('No recommendations available.')

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_pdf_report(
    patient_info: Dict[str, Any],
    recommendations: Dict[str, Any],
    big5_scores: Dict[str, Any],
) -> bytes:
    from fpdf import FPDF  # type: ignore
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font('Helvetica', 'B', 16)
    pdf.cell(0, 10, 'TheraMuse Recommendation Report', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 11)
    pdf.cell(0, 8, f'Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}', new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font('Helvetica', 'B', 13)
    pdf.cell(0, 8, 'Patient Overview', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 11)
    for label, value in _patient_rows(patient_info):
        pdf.multi_cell(170, 6, f'{label}: {_sanitize_for_pdf(value)}', new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(2)
    pdf.set_font('Helvetica', 'B', 13)
    pdf.cell(0, 8, 'Big Five Scores', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 11)
    for trait, score in _big5_rows(big5_scores):
        pdf.cell(0, 6, f'{trait}: {_sanitize_for_pdf(score)}', new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.ln(2)
    pdf.set_font('Helvetica', 'B', 13)
    pdf.cell(0, 8, 'Recommendations', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 11)
    rec_rows = _recommendation_rows(recommendations)
    if rec_rows:
        current_category = None
        for row in rec_rows:
            if row['category'] != current_category:
                pdf.ln(2)
                pdf.set_font('Helvetica', 'B', 12)
                pdf.multi_cell(170, 6, _sanitize_for_pdf(row['category']), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font('Helvetica', '', 11)
                current_category = row['category']
            # Format the title and metadata properly
            title_text = f"{row['rank']}. {_sanitize_for_pdf(row['title'])}"
            pdf.multi_cell(170, 6, title_text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            # Artist info if available
            if isinstance(row, dict) and 'artist' in row and row['artist']:
                artist_info = f"Artist: {_sanitize_for_pdf(row['artist'])}"
                pdf.multi_cell(170, 5, artist_info, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            # Channel and URL info on separate lines
            channel_info = f"Channel: {_sanitize_for_pdf(row.get('channel') or 'N/A')}"
            pdf.multi_cell(170, 5, channel_info, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            if row.get('url'):
                url_info = f"URL: {_sanitize_for_pdf(row['url'])}"
                pdf.multi_cell(170, 5, url_info, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

            if row.get('description'):
                pdf.multi_cell(170, 5, f"Notes: {_sanitize_for_pdf(row['description'])}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(2)  # Add spacing between songs
    else:
        pdf.multi_cell(170, 6, 'No recommendations available.', new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    return pdf.output()






def _patient_rows(patient_info: Dict[str, Any]) -> Iterable[Tuple[str, str]]:
    if not patient_info:
        return (('Details', 'Not provided'),)

    preferred_order = [
        'name',
        'age',
        'sex',
        'condition',
        'birthplaceCity',
        'birthplaceCountry',
    ]
    rows: List[Tuple[str, str]] = []

    for key in preferred_order:
        if key in patient_info:
            rows.append((_labelize(key), _stringify(patient_info[key])))

    for key, value in patient_info.items():
        if key not in preferred_order:
            rows.append((_labelize(key), _stringify(value)))

    return rows


def _big5_rows(scores: Dict[str, Any]) -> Iterable[Tuple[str, str]]:
    if not scores:
        return (('Scores', 'Not provided'),)
    order = [
        ('openness', 'Openness'),
        ('conscientiousness', 'Conscientiousness'),
        ('extraversion', 'Extraversion'),
        ('agreeableness', 'Agreeableness'),
        ('neuroticism', 'Neuroticism'),
    ]
    rows: List[Tuple[str, str]] = []
    for key, label in order:
        if key in scores:
            rows.append((label, _stringify(scores[key])))
    for key, value in scores.items():
        if key not in dict(order):
            rows.append((_labelize(key), _stringify(value)))
    return rows


def _recommendation_rows(recommendations: Dict[str, Any]) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    categories = recommendations.get('categories') if isinstance(recommendations, dict) else None
    if not isinstance(categories, dict):
        return rows
    for key, payload in categories.items():
        songs = payload.get('songs') if isinstance(payload, dict) else None
        if not isinstance(songs, list):
            continue
        label = _labelize(payload.get('label') or key)
        for idx, song in enumerate(songs, start=1):
            if isinstance(song, dict):
                rows.append(
                    {
                        'category': label,
                        'rank': idx,
                        'title': _stringify(song.get('title', '')),
                        'artist': _stringify(song.get('artist', '')),
                        'channel': _stringify(song.get('channel', '')),
                        'description': _stringify(song.get('description', '')),
                        'url': _stringify(song.get('url', '')),
                    }
                )
            else:
                # Handle simple string titles
                rows.append(
                    {
                        'category': label,
                        'rank': idx,
                        'title': _stringify(song),
                        'artist': '',
                        'channel': '',
                        'description': '',
                        'url': '',
                    }
                )
    return rows


def _stringify(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, (list, tuple, set)):
        return ', '.join(_stringify(item) for item in value if item is not None)
    if isinstance(value, dict):
        return json.dumps(_make_json_safe(value), ensure_ascii=False)
    return str(value)


def _labelize(raw: Any) -> str:
    text = _stringify(raw)
    if not text:
        return ''
    return text.replace('_', ' ').replace('-', ' ').strip().title()


def _sanitize_for_pdf(text: str) -> str:
    """Replace Unicode characters with ASCII equivalents for PDF compatibility."""
    if not isinstance(text, str):
        text = str(text)

    # Common Unicode character replacements
    replacements = {
        '\u2022': '*',  # Bullet point
        '\u2026': '...',  # Ellipsis
        '\u201c': '"',  # Left double quote
        '\u201d': '"',  # Right double quote
        '\u2018': "'",  # Left single quote
        '\u2019': "'",  # Right single quote
        '\u2013': '-',  # En dash
        '\u2014': '--', # Em dash
        '\u00a9': '(c)', # Copyright
        '\u00ae': '(r)', # Registered trademark
        '\u2122': '(tm)', # Trademark
        '\u00b0': 'deg', # Degree symbol
        '\u00b1': '+/-', # Plus-minus
        '\u00d7': 'x',   # Multiplication sign
        '\u00f7': '/',   # Division sign
        '\u2264': '<=',  # Less than or equal
        '\u2265': '>=',  # Greater than or equal
        '\u2260': '!=',  # Not equal
        '\u2212': '-',   # Minus sign
        '\u00a0': ' ',   # Non-breaking space
    }

    # Replace each Unicode character with its ASCII equivalent
    for unicode_char, ascii_char in replacements.items():
        text = text.replace(unicode_char, ascii_char)

    # Remove any remaining non-ASCII characters
    return text.encode('ascii', 'ignore').decode('ascii')


def _make_json_safe(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {str(key): _make_json_safe(value) for key, value in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [_make_json_safe(item) for item in obj]
    if isinstance(obj, datetime):
        return obj.isoformat()
    return obj

